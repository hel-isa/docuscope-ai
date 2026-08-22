import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pypdf import PdfWriter

from app.config import MAX_IMAGE_PIXELS
from app.parsers.docx_parser import parse_docx
from app.parsers.image_parser import parse_image
from app.parsers.pdf_parser import parse_pdf
from app.parsers.xlsx_parser import parse_xlsx
from app.security.resource_guards import (
    ImageTooLargeError,
    PdfTooManyPagesError,
    ZipBombSuspectedError,
    configure_pillow_limits,
)


def test_parse_docx_rejects_zip_bomb(tmp_path, monkeypatch) -> None:
    f = tmp_path / "bomb.docx"
    doc = Document()
    doc.add_paragraph("This paragraph repeats to build up uncompressed size. " * 200)
    doc.save(str(f))

    monkeypatch.setattr("app.parsers.docx_parser.MAX_ZIP_UNCOMPRESSED_BYTES", 100)

    with pytest.raises(ZipBombSuspectedError):
        parse_docx(f)


def test_parse_xlsx_rejects_zip_bomb(tmp_path, monkeypatch) -> None:
    f = tmp_path / "bomb.xlsx"
    wb = Workbook()
    ws = wb.active
    for row in range(500):
        ws.append([f"repeated value {row}"] * 10)
    wb.save(str(f))

    monkeypatch.setattr("app.parsers.xlsx_parser.MAX_ZIP_UNCOMPRESSED_BYTES", 100)

    with pytest.raises(ZipBombSuspectedError):
        parse_xlsx(f)


def test_parse_pdf_rejects_too_many_pages(tmp_path, monkeypatch) -> None:
    writer = PdfWriter()
    for _ in range(20):
        writer.add_blank_page(width=72, height=72)

    f = tmp_path / "many_pages.pdf"
    with f.open("wb") as out:
        writer.write(out)

    monkeypatch.setattr("app.parsers.pdf_parser.MAX_PDF_PAGES", 5)

    with pytest.raises(PdfTooManyPagesError):
        parse_pdf(f)


class _FakePage:
    def extract_text(self) -> str:
        return "page text"


class _UnreliableCountPages(list):
    """
    Simulates a malformed page tree where accessing len() raises (e.g. a
    corrupted /Count entry) but iteration still yields real pages. This is
    exactly the shape an attacker would craft to bypass an upfront-only
    page-count guard.
    """

    def __len__(self) -> int:
        raise ValueError("simulated corrupt page tree /Count")


def test_parse_pdf_rejects_too_many_pages_even_when_upfront_count_fails(
    tmp_path, monkeypatch
) -> None:
    f = tmp_path / "unreliable_count.pdf"
    f.write_bytes(b"%PDF-1.4\n")

    class _FakeReader:
        def __init__(self, _path: str) -> None:
            self.pages = _UnreliableCountPages([_FakePage() for _ in range(20)])
            self.metadata = {}

    monkeypatch.setattr("app.parsers.pdf_parser.PdfReader", _FakeReader)
    monkeypatch.setattr("app.parsers.pdf_parser.MAX_PDF_PAGES", 5)

    with pytest.raises(PdfTooManyPagesError):
        parse_pdf(f)


def test_parse_image_rejects_decompression_bomb(tmp_path) -> None:
    f = tmp_path / "normal.png"
    Image.new("RGB", (200, 200), color=(0, 128, 255)).save(f)

    configure_pillow_limits(max_pixels=1000)  # tiny limit for this test
    try:
        with pytest.raises(ImageTooLargeError):
            parse_image(f)
    finally:
        configure_pillow_limits(max_pixels=MAX_IMAGE_PIXELS)  # restore real default
